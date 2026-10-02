#!/usr/bin/env python3
"""
organic_max_sheet.py — writes the SCALE family (R30–R36, "Scale plan to $250K", tools/organic_engine.py sc_run; BLITZ.md
§14) into economics.xlsx sheet `Organic_Max` (rebuilt from scratch; every other sheet is left untouched and verified,
Projection_Aggressive kept) and appends r30…r36 *_booked_MRR / *_retained_MRR / *_cash_scale / *_views_scale /
*_posts_scale columns to mrr_blitz_daily.csv (the 69 existing columns, incl. the frozen legacy Organic-max r30_MRR…
r35_posts, are left byte-identical).

    python3 tools/organic_max_sheet.py

Sheet layout: definitions, inputs block (row 6+), ladder, run table, sensitivity, 3-way comparison, R30A check against
data/projection_aggressive_central.csv; daily block from row DAILY0 (day, then booked / retained / cash / views / posts
per run).
"""
import csv, hashlib, os, re, sys
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter as L
import organic_engine as oe

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
XLSX = os.path.join(ROOT, "economics.xlsx")
CSV = os.path.join(ROOT, "mrr_blitz_daily.csv")
SHEET = "Organic_Max"
DAILY0 = 200           # header row of the daily block; day 1 on DAILY0 + 1
DAYS = 360             # csv horizon (the approved projection is 180 days; days 181–360 are model extrapolation)
SERIES = [("booked_MRR", "booked_MRR"), ("retained_MRR", "retained_MRR"), ("cash_scale", "cash"),
          ("views_scale", "views_day"), ("posts_scale", "posts_day")]
CHECK3 = ["Projection_Aggressive", "Organic_First", "Costs"]
B = Font(bold=True)
HEAD = PatternFill("solid", fgColor="000000"); HF = Font(bold=True, color="FFFFFF")
WRAP = Alignment(wrap_text=True, vertical="top")


def cells(ws):
    return {c.coordinate: c.value for row in ws.iter_rows() for c in row if c.value is not None}


def snapshot(wb, skip):
    h = {}
    for ws in wb.worksheets:
        if ws.title == skip:
            continue
        m = hashlib.md5()
        for row in ws.iter_rows():
            for c in row:
                if c.value is not None:
                    m.update(f"{c.coordinate}={c.value!r};".encode())
        h[ws.title] = (m.hexdigest(), ws.max_row, ws.max_column)
    return h


def hdr(ws, r, vals):
    for j, v in enumerate(vals, 1):
        c = ws.cell(r, j, v); c.font = HF; c.fill = HEAD; c.alignment = WRAP


def put(ws, r, vals):
    for j, v in enumerate(vals, 1):
        ws.cell(r, j, round(v, 2) if isinstance(v, float) else v)


def feature_count():
    s = open(os.path.join(ROOT, "SYSTEM_RECAP.md")).read()
    sec = s.split("## 2. Features by layer")[1].split("\n## 3.")[0]
    rows = [x for x in sec.splitlines() if x.startswith("|") and not re.match(r"^\|\s*-", x)]
    return len([x for x in rows if not x.lower().startswith("| feature")])


def compute():
    runs = [(r, oe.sc_run(r, DAYS)) for r in oe.RUNS_SCALE]
    r30a = oe.sc_run(oe.R30A, 180)
    sums = [(r, rows, oe.sc_summary(rows[:180])) for r, rows in runs]
    sens = [(nm, oe.sc_summary(oe.sc_run(dict(oe.R30, **ch), 180))) for nm, ch in oe.SENS_SCALE]
    orig_rows, osm = oe.original_ym()
    rows20, pre20, info20 = oe.run(oe.RUNS[0]); s20 = oe.summ(rows20, pre20, info20)
    return dict(runs=runs, sums=sums, r30a=r30a, sens=sens, orig=(orig_rows, osm), r20=(rows20, s20),
                verify=oe.sc_verify_csv(), features=feature_count())


def comparison(out):
    o_rows, osm = out["orig"]; rows20, s20 = out["r20"]; r30 = out["sums"][0][2]; rows30 = out["runs"][0][1]
    og = lambda d, k: o_rows[d - 1][k]; g20 = lambda d, k: rows20[d - 1][k]; g30 = lambda d, k: rows30[d - 1][k]
    opex = (oe.SH["fixed"] + oe.SH["shop_plan"]) / 30
    mg = lambda cost, mrr: 100 * (1 - 30 * cost / mrr) if mrr > 0 else None
    nf = out["features"]
    return [
        ("posts/day (d30)", 3, g20(30, "posts"), g30(30, "posts_day")),
        ("views/day d30", og(30, "views"), g20(30, "views"), g30(30, "views_day")),
        ("views/day d90", og(90, "views"), g20(90, "views"), g30(90, "views_day")),
        ("retained MRR d30", og(30, "mrr"), g20(30, "ret"), g30(30, "retained_MRR")),
        ("retained MRR d60", og(60, "mrr"), g20(60, "ret"), g30(60, "retained_MRR")),
        ("retained MRR d90", og(90, "mrr"), g20(90, "ret"), g30(90, "retained_MRR")),
        ("retained MRR d180", og(180, "mrr"), g20(180, "ret"), g30(180, "retained_MRR")),
        ("cost/day d30 (model basis)", og(30, "cost"), g20(30, "costs"), g30(30, "cost_day")),
        ("cost/day d30 excluding team opex", og(30, "cost"), g20(30, "costs") - opex, g30(30, "cost_day")),
        ("margin d30 % = 1 − 30 × cost/day ÷ MRR (booked; ORIGINAL paid MRR)", mg(og(30, "cost"), og(30, "mrr")),
         mg(g20(30, "costs"), g20(30, "mrr")), r30["mg30"]),
        ("margin d90 %", mg(og(90, "cost"), og(90, "mrr")), mg(g20(90, "costs"), g20(90, "mrr")), g30(90, "margin_pct")),
        ("breakeven day (cumulative cash ≥ 0)", osm["be"], s20["beday"], r30["be"]),
        ("features (SYSTEM_RECAP §2 rows; ORIGINAL observed; UPDATED = + spend gate/cap + scale ladder)", 8, nf, nf + 2),
    ]


def main():
    out = compute()
    runs, sums = out["runs"], out["sums"]

    # ---------------- csv ----------------
    with open(CSV, newline="") as f:
        rd = csv.DictReader(f); fields = list(rd.fieldnames); data = list(rd)
    new = [f"{t}_{k}" for t in oe.SC_TAGS for k, _ in SERIES]
    fields += [c for c in new if c not in fields]
    byday = {int(float(x["day"])): x for x in data}
    for tag, (r, rows) in zip(oe.SC_TAGS, runs):
        for w in rows:
            x = byday.get(w["day"])
            if x is None:
                continue
            for k, src in SERIES:
                x[f"{tag}_{k}"] = f"{w[src]:.2f}" if k in ("booked_MRR", "retained_MRR", "cash_scale") else f"{w[src]:.0f}"
    with open(CSV, "w", newline="") as f:
        w = csv.DictWriter(f, fieldnames=fields); w.writeheader(); w.writerows(data)

    # ---------------- workbook ----------------
    wb = openpyxl.load_workbook(XLSX)
    before = snapshot(wb, SHEET)
    before3 = {n: cells(wb[n]) for n in CHECK3}
    assert "Projection_Aggressive" in wb.sheetnames
    if SHEET in wb.sheetnames:
        del wb[SHEET]
    ws = wb.create_sheet(SHEET)
    ws.column_dimensions["A"].width = 48; ws.column_dimensions["B"].width = 62
    for j in range(3, 40):
        ws.column_dimensions[L(j)].width = 12
    ws["A1"] = "Scale plan to $250K — runs R30–R36 (engine tools/organic_engine.py sc_run; writer tools/organic_max_sheet.py; BLITZ.md §14)"
    ws["A1"].font = Font(bold=True, size=13)
    ws["A2"] = ("Port of the client-approved projection (data/projection_aggressive_central.csv, sheet Projection_Aggressive). "
                "booked_MRR = 0.95 × ($25 × active members + $147 × coached). retained_MRR = 0.95 × $25 × (50% of first-cycle members "
                "+ every renewed member): PLAN ON THIS ONE. cash = −$15K + $12 per buyer + $25 per renewal + coached $147/30 per day − 3% fees "
                "− daily cost (no team opex). contracted_30d = renewals due in 30 days from survivors = retained_MRR (monthly billing). "
                "R30–R36 step the ladder, open the $30K paid gate and size the 25% all-in cap on TRAILING-7-DAY RETAINED MRR "
                "(= workers/growth/governor.py); R30A is the approved file exactly (it steps them on the previous day's BOOKED MRR).")
    ws["A2"].alignment = WRAP; ws.row_dimensions[2].height = 75
    r = 4
    ws.cell(r, 1, "1. Inputs (central, range, source). Every value without a citation is an ASSUMPTION.").font = B; r += 1
    hdr(ws, r, ["key", "input", "central", "low", "high", "unit", "source / rationale"]); r += 1
    for key, label, c, lo, hi, unit, src in oe.SC_INPUTS:
        put(ws, r, [key, label, c, lo, hi, unit, src]); ws.cell(r, 7).alignment = WRAP; r += 1
    r += 1
    ws.cell(r, 1, "2. Scale-on-MRR ladder (= workers/growth/config.py governor.scale_rules; never steps down)").font = B; r += 1
    hdr(ws, r, ["trailing-7-day retained MRR ≥", "pages open", "masters/page/day", "Trial Reels/page/day", "generation tier"]); r += 1
    for row in oe.SC_LADDER:
        put(ws, r, list(row)); r += 1
    r += 1

    ws.cell(r, 1, "3. Run table (days 1–180)").font = B; r += 1
    ms = oe.SC_MILESTONES
    cols = (["run", "definition"] + [f"booked ${x // 1000}K day" for x in ms] + [f"retained ${x // 1000}K day" for x in ms]
            + [f"booked d{d}" for d in (30, 60, 90, 180)] + [f"retained d{d}" for d in (30, 60, 90, 180)]
            + ["cash low", "cash low day", "breakeven day", "margin d30 %", "margin d60 %", "gate open day",
               "ladder $10K day", "ladder $30K day", "ladder $50K day", "ladder $100K day", "page ceiling first binds", "all pages at ceiling",
               "posts/day d30", "views/day d30", "views/day d90", "cost/day d30"])
    hdr(ws, r, cols); r += 1
    allsum = sums + [(oe.R30A, out["r30a"], oe.sc_summary(out["r30a"]))]
    for (rr, rows, sm) in allsum:
        tag = rr["name"].split()[0]
        put(ws, r, [tag, rr["name"]] + [sm[f"bk{x // 1000}K"] for x in ms] + [sm[f"rt{x // 1000}K"] for x in ms]
            + [sm[f"bk{d}"] for d in (30, 60, 90, 180)] + [sm[f"rt{d}"] for d in (30, 60, 90, 180)]
            + [sm["low"], sm["lowday"], sm["be"], sm["mg30"], sm["mg60"], sm["gate"]] + [sm["ladder"][x] for x in (10000, 30000, 50000, 100000)]
            + [sm["ceil1"], sm["ceil_all"], sm["p30"], sm["v30"], sm["v90"], sm["c30"]])
        r += 1
    ws.cell(r, 1, "Blank milestone = not reached within 180 days. Ladder day = first day the row is in effect.").alignment = WRAP
    r += 2

    ws.cell(r, 1, "4. Sensitivity on R30 (one input at a time)").font = B; r += 1
    hdr(ws, r, ["change", "", "retained d30", "retained d60", "retained d90", "retained d180", "retained $100K day", "retained $250K day",
                "booked $250K day", "gate open day", "all pages at ceiling"]); r += 1
    s30 = sums[0][2]
    for nm, sm in [("R30 base", s30)] + out["sens"]:
        put(ws, r, [nm, "", sm["rt30"], sm["rt60"], sm["rt90"], sm["rt180"], sm["rt100K"], sm["rt250K"], sm["bk250K"], sm["gate"], sm["ceil_all"]])
        r += 1
    r += 1

    ws.cell(r, 1, "5. 3-way comparison").font = B; r += 1
    hdr(ws, r, ["measure", "", "ORIGINAL (YM-style: 1 page, 3/day, $19.99 ebook + Whop $19.99/mo; POSTDB)", "CURRENT (R20)", "UPDATED (R30)"]); r += 1
    for row in comparison(out):
        put(ws, r, [row[0], ""] + list(row[1:])); r += 1
    ws.cell(r, 1, ("ORIGINAL retained = paid Whop members after the 3-day trial (S curve); cost 3 posts × $1.12 + $16/day, no opex. "
                   "CURRENT cost includes $30.5K/month team opex + Shopify plan ($1,018/day); UPDATED (projection basis) carries no team opex."))
    ws.cell(r, 1).alignment = WRAP
    r += 2

    ws.cell(r, 1, "6. R30A vs data/projection_aggressive_central.csv (max abs error, max relative error, day)").font = B; r += 1
    hdr(ws, r, ["column", "", "max abs error", "relative", "day"]); r += 1
    for c, (err, rel, day) in out["verify"].items():
        put(ws, r, [c, "", err, rel, day]); r += 1
    assert r < DAILY0 - 2, f"layout overflow: row {r} reaches the daily block"

    ws.cell(DAILY0 - 1, 1, f"Daily rows (days 1–{DAYS}; days 181–{DAYS} extrapolate beyond the approved 180-day file)").font = B
    hdr(ws, DAILY0, ["day"] + [f"{t} {k}" for t in oe.SC_TAGS for k, _ in SERIES])
    for i in range(DAYS):
        vals = [i + 1]
        for _, rows in runs:
            w = rows[i]
            vals += [round(w[src], 2) for _, src in SERIES]
        for j, v in enumerate(vals, 1):
            ws.cell(DAILY0 + 1 + i, j, v)
    wb.save(XLSX)

    # ---------------- verify ----------------
    wb2 = openpyxl.load_workbook(XLSX)
    after = snapshot(wb2, SHEET)
    assert before == after, "another sheet changed"
    for n in CHECK3:
        assert cells(wb2[n]) == before3[n], f"{n} changed"
    ws2 = wb2[SHEET]
    for i in (0, 29, 179, DAYS - 1):
        for k, (_, rows) in enumerate(runs):
            for m, (_, src) in enumerate(SERIES):
                assert ws2.cell(DAILY0 + 1 + i, 2 + k * len(SERIES) + m).value == round(rows[i][src], 2)
    print(f"Organic_Max rebuilt; {len(after)} other sheets hash-identical; cell-for-cell identical: {', '.join(CHECK3)}; "
          f"csv +{len(new)} columns")
    return out


if __name__ == "__main__":
    main()
