#!/usr/bin/env python3
"""
organic_sheet_refresh.py — writes the organic_engine.py results into economics.xlsx sheet `Organic_First`
(values only; the sheet's result formulas over the daily rows stay live) and refreshes the r20…r26 columns
of mrr_blitz_daily.csv. Run after any change to tools/organic_engine.py:

    python3 tools/organic_engine.py --json /tmp/organic.json   # optional; this script re-runs the engine itself
    python3 tools/organic_sheet_refresh.py

Layout contract (the sheet was laid out by the original builder; this script only overwrites cells in place):
  rows 85–120   per-run inputs, columns B..H = R20..R26
  rows 124/126–129/161–163   result values that are not formulas
  rows 167–175  runway table (3 cases × 7/14/21)
  rows 180–187  closest-at-least-cash combos (cell × cap/shoutouts)
  rows 191+     sensitivity (R20 then R23 bases), one row per SENS entry
  rows 294+     configuration ladder (3 scenarios × 4 caps × 3 shoutout levels)
  row 360 + i*378   daily block i (title row, header row, 14 runway rows, 360 day rows)
"""
import csv, os, sys
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import openpyxl
import organic_engine as oe

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
XLSX = os.path.join(ROOT, "economics.xlsx")
CSV = os.path.join(ROOT, "mrr_blitz_daily.csv")

INPUT_ROWS = [  # sheet row → run key (rows 85..120)
    (85, "name"), (86, "cell"), (87, "fe_price"), (88, "runway"), (89, "scen"), (90, "posts_max"), (91, "wl_rate"), (92, "c72"), (93, "ctail"),
    (94, "seed_wl"), (95, "cvr_org"), (96, "cvr_cold"), (97, "cvr_warm"), (98, "elig_addr"), (99, "elig_pay"), (100, "pp_take"), (101, "later_take"),
    (102, "subB"), (103, "renB1"), (104, "ref_pp"), (105, "price"), (106, "lpv"), (107, "cpm_mult"), (108, "exist"), (109, "unig"), (110, "k9"),
    (111, "aff"), (112, "sh_max"), (113, "sh_ctr"), (114, "cap"), (115, "boost_eff"), (116, "rt_eff"), (117, "gate_on"), (118, "gate_sp"),
    (119, "ann_take"), (120, "save_take"),
]
DAY_KEYS = ["d", "pages", "posts", "views", "visitors", "v_org", "wl_buy", "v_warm", "v_sh", "v_aff", "sp_cold", "sp_boost", "sp_rt", "sp", "v_paid", "shp",
            "orders", "E", "P", "accept", "later", "B", "pB", "refB", "fe_rev", "members", "mrr", "ret", "gross", "receipts", "costs", "net", "cash", "res",
            "cum_media", "cum_fe", "cac"]  # 37 value columns; column 38 is the sheet's own formula
BLOCK = 378
FIRST_BLOCK_TITLE = 360


def main():
    res, rw, sens, ladder = oe.main()
    wb = openpyxl.load_workbook(XLSX)
    ws = wb["Organic_First"]
    assert str(ws.cell(85, 1).value).startswith("Run") and str(ws.cell(361, 1).value) == "Day", "Organic_First layout changed; update this script"

    ws.cell(1, 1).value = ("Organic_First — organic-first launch on the CANON UPDATE 2/3 Shopify ladder (runs R20–R26): LAUNCH DEFAULT cell B "
                           '"$12 today = books + first month, then $25/mo" (founding plan + STARTER12 first-payment code on the free Shopify Subscriptions app); '
                           "TEST CELL A = $12 one-time books → founding offer on the thank-you page + 3 emails (one-click post-purchase OFF); "
                           "runway 7/14/21; waitlist / list solver; amplification-only paid; §11 gate. Engine: tools/organic_engine.py (values); "
                           "result formulas are live over the daily rows. Refreshed by tools/organic_sheet_refresh.py.")

    # 2. per-run inputs
    for ci, (r, rows, pre, info, sm) in enumerate(res, start=2):
        ws.cell(84, ci).value = r["name"][:3]
        for row, key in INPUT_ROWS:
            ws.cell(row, ci).value = r[key]
        # 3. result values that are not formulas
        ws.cell(124, ci).value = info["W0"]
        ws.cell(126, ci).value = info["mpo"]
        ws.cell(127, ci).value = info["fe_unit"]
        ws.cell(128, ci).value = info["contribA"]
        ws.cell(129, ci).value = info["contribB"]
        ws.cell(161, ci).value = info["gate"]
        ws.cell(162, ci).value = info["gate_cost"]
        ws.cell(163, ci).value = info["gate_line"]
        # daily block
        t = FIRST_BLOCK_TITLE + (ci - 2) * BLOCK
        ws.cell(t, 1).value = f"Daily rows — {r['name']}"
        for i, p in enumerate(pre):
            rr = t + 2 + i
            ws.cell(rr, 1).value = p["d"]; ws.cell(rr, 2).value = p["pages"]; ws.cell(rr, 3).value = p["posts"]; ws.cell(rr, 4).value = p["views"]
            ws.cell(rr, 5).value = p["visitors"]; ws.cell(rr, 7).value = p["waitlist"]; ws.cell(rr, 33).value = p["cash"]
        for i, d in enumerate(rows):
            rr = t + 16 + i
            for cj, k in enumerate(DAY_KEYS, start=1):
                ws.cell(rr, cj).value = d[k]

    # 4. runway table
    for i, x in enumerate(rw):
        rr = 167 + i
        sm = x["summ"]
        vals = [x["name"], x["R"], sm["wl_org"], -sm["runway_cost"], sm["m4"], sm["m14"], sm["m30"],
                x["need"]["10000@4"], x["need"]["50000@14"], x["need"]["100000@30"],
                x["need_list"]["10000@4"], x["need_list"]["50000@14"], x["need_list"]["100000@30"]]
        for cj, v in enumerate(vals, start=1):
            ws.cell(rr, cj).value = v
    # 4b. combos (the engine's combo list is rebuilt here in the sheet's order: cell × tier)
    combos = []
    for cap, sh in ((0, 0), (500, 2), (1500, 2), (3000, 4)):
        for cell in ("B", "A"):
            base = dict(oe.RUNS[0] if cell == "B" else oe.RUNS[2], cap=cap, sh_max=sh)
            row = [cell, cap, sh]
            media = None
            for tgt, day in oe.MILESTONES:
                seed = oe.solve_seed(base, tgt, day)
                if seed is None:
                    row += [None, None, None]
                    continue
                rows_, pre_, info_ = oe.run(dict(base, seed_wl=seed)); sm_ = oe.summ(rows_, pre_, info_)
                row += [seed, sm_["low"], sm_["ret30"]]
                media = sm_["media30"]
            row.append(media)
            combos.append(row)
    for i, row in enumerate(combos):
        for cj, v in enumerate(row, start=1):
            ws.cell(180 + i, cj).value = v
    # 5. sensitivity
    rr = 191
    for x in sens:
        sm = x["summ"]
        vals = [x["base"], x["name"], sm["m4"], sm["m14"], sm["m30"], x["d30_delta"], sm["ret30"], sm["m90"], sm["m180"], sm["low"], sm["d10"], sm["d50"], sm["d100"], "FLAG" if x["flag"] else None]
        for cj, v in enumerate(vals, start=1):
            ws.cell(rr, cj).value = v
        rr += 1
    # clear any leftover sensitivity rows up to the ladder header (row 292)
    while rr < 292 and ws.cell(rr, 1).value not in (None, ""):
        for cj in range(1, 15):
            ws.cell(rr, cj).value = None
        rr += 1
    # 6. ladder
    for i, x in enumerate(ladder):
        sm = x["summ"]
        vals = [x["scen"], x["cap"], x["sh"], sm["m4"], sm["m14"], sm["m30"], sm["ret30"], sm["m90"], sm["m180"], sm["d10"], sm["d50"], sm["d100"], sm["low"], sm["media30"], sm["fe30"]]
        for cj, v in enumerate(vals, start=1):
            ws.cell(294 + i, cj).value = v
    wb.save(XLSX)

    # csv: r20…r26 MRR and cash columns
    with open(CSV, newline="") as f:
        rd = csv.DictReader(f); fields = rd.fieldnames; data = list(rd)
    byday = {int(float(x["day"])): x for x in data}
    for ci, (r, rows, pre, info, sm) in enumerate(res):
        tag = f"r{20 + ci}"
        for d in rows:
            x = byday.get(d["d"])
            if x is None:
                continue
            if f"{tag}_MRR" in fields: x[f"{tag}_MRR"] = f"{d['mrr']:.2f}"
            if f"{tag}_cash" in fields: x[f"{tag}_cash"] = f"{d['cash']:.2f}"
    with open(CSV, "w", newline="") as f:
        w = csv.DictWriter(f, fieldnames=fields); w.writeheader(); w.writerows(data)
    print("refreshed", XLSX, "and", CSV)


if __name__ == "__main__":
    main()
