"""Content-production cost model -> data/costs_model.csv (plain numbers) and the formula-driven `Costs` sheet in
economics.xlsx (inputs block + scenario grid; other sheets untouched). Prices: public list prices checked Oct 1 2026
(sources in COSTS.md and in the sheet). [A] = assumption/estimate. No network, no spend.
"""
from __future__ import annotations

import csv
import json
import math
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
OUT_CSV = ROOT / "data/costs_model.csv"
XLSX = ROOT / "economics.xlsx"
DPM = 30.4

# (key, value, unit, source/label)
INPUTS = [
    ("p_img", 0.067, "$/image", "Nano Banana 2 (gemini-3.1-flash-image) 1K, standard: ai.google.dev/gemini-api/docs/pricing (batch $0.034)"),
    ("p_ls", 0.0562, "$/s", "fal Kling AI Avatar v2 Standard (lip-sync from still): fal.ai/models/fal-ai/kling-video/ai-avatar/v2/standard"),
    ("p_mv", 0.07, "$/s", "fal Kling 2.6 Standard Motion Control: fal.ai/models/fal-ai/kling-video/v2.6/standard/motion-control"),
    ("p_veo", 0.15, "$/s", "Veo 3.1 Fast i2v on fal, no audio: costgoat.com/pricing/google-veo (Gemini API direct $0.10/s)"),
    ("p_sin", 2.0, "$/MTok", "Claude Sonnet 5.5 input: platform.claude.com/docs/en/about-claude/pricing"),
    ("p_sout", 10.0, "$/MTok", "Claude Sonnet 5.5 output: same"),
    ("p_hin", 1.0, "$/MTok", "Claude Haiku 4.5 input: same"),
    ("p_hout", 5.0, "$/MTok", "Claude Haiku 4.5 output: same"),
    ("el_ovg", 0.10, "$/1K chars", "ElevenLabs overage, Multilingual v2/v3: flexprice.io/blog/elevenlabs-pricing-breakdown (API PAYG v3 $0.08: elevenlabs.io/pricing/api)"),
    ("el_creator_fee", 22, "$/mo", "ElevenLabs Creator, 121K credits (1 credit ≈ 1 char)"),
    ("el_creator_inc", 121000, "chars/mo", "same"),
    ("el_pro_fee", 99, "$/mo", "ElevenLabs Pro, 600K credits"),
    ("el_pro_inc", 600000, "chars/mo", "same"),
    ("el_scale_fee", 299, "$/mo", "ElevenLabs Scale, 1.8M credits"),
    ("el_scale_inc", 1800000, "chars/mo", "same"),
    ("el_bus_fee", 990, "$/mo", "ElevenLabs Business, 6M credits"),
    ("el_bus_inc", 6000000, "chars/mo", "same"),
    ("review_rate", 20.0, "$/h", "[A] contract compliance reviewer"),
    ("review_s", 20, "s/post", "brief: 20 s human review per judged post"),
    ("shoot_oneoff", 2250, "$", "[A] performer shoot, midpoint of PIPELINE.md $1.5-3K"),
    ("shoot_refresh_q", 500, "$/quarter", "PIPELINE.md performer refresh"),
    ("shoot_amort_days", 90, "days", "[A] amortization window"),
    ("vps_box", 90, "$/mo", "[A] 8 dedicated vCPU box (Hetzner CCX33 class after the Jun 15 2026 price adjustment: docs.hetzner.com/general/infrastructure-and-availability/price-adjustment/)"),
    ("assembly_min_file", 1.5, "CPU-min/file", "[A] ffmpeg packaging per video file"),
    ("box_cpu_min_day", 5760, "CPU-min/day", "[A] 8 vCPU x 20 h x 60% utilisation"),
    ("r2_gb", 0.015, "$/GB-mo", "Cloudflare R2 standard, egress free: developers.cloudflare.com/r2/pricing/"),
    ("gb_file", 0.025, "GB/file", "[A] 45 s 1080x1920 at ~4.5 Mbps"),
    ("gb_master_tmp", 0.05, "GB/master", "[A] intermediates kept"),
    ("retain_days", 90, "days", "[A] retention"),
    ("supabase", 40, "$/mo", "Supabase Pro $25 (supabase.com/pricing) + [A] $15 compute upgrade"),
    ("vercel", 40, "$/mo", "Vercel Pro $20/seat x 2 seats [A]: vercel.com/pricing"),
    ("n8n", 0, "$/mo", "n8n Community self-hosted, no licence fee (runs on the VPS): nocode.mba/articles/n8n-pricing"),
    ("email", 20, "$/mo", "Resend Pro 50K emails: resend.com/docs/knowledge-base/what-is-resend-pricing"),
    ("manychat_page", 29, "$/mo/page", "ManyChat Pro (2,500 active contacts, then $0.05): manychat.com/pricing"),
    ("posting_api", 0, "$/mo", "Direct platform APIs from n8n (no scheduler SaaS)"),
    ("sec", 42, "s", "[A] mean master length (allocator bucket M target; library mean ~42 s)"),
    ("swap_s", 8, "s", "2 last-line CTA swaps (TikTok, YouTube) x 4 s re-lip-sync"),
    ("swap_chars", 160, "chars", "2 x 80 chars re-TTS"),
    ("chars_s", 15, "chars/s", "[A] spoken ~2.6 words/s"),
    ("rt_img", 1.5, "x", "[A] image retakes"), ("rt_ls", 1.15, "x", "[A] lip-sync retakes"),
    ("rt_mv", 1.3, "x", "[A] motion-control retakes"), ("rt_ins", 1.25, "x", "[A] insert retakes"), ("rt_tts", 1.3, "x", "[A] TTS retakes"),
    ("img_th", 3, "images", "render_day1.py: 1 keyframe + 2 insert stills"), ("ins_th", 16, "s", "render_day1.py: 2 Veo inserts x 8 s"),
    ("img_mv", 3, "images", "[A] 2 refs + 1 insert still"), ("mv_share", 0.6, "of runtime", "[A] motion-control share of a movement master"),
    ("ins_mv", 8, "s", "[A] 1 insert"), ("img_in", 7, "images", "[A] 3 clip stills + 4 Ken Burns stills"), ("ins_in", 24, "s", "[A] 3 Veo inserts x 8 s"),
    ("carousel_share", 0.3, "of text posts", "[A]"), ("carousel_imgs", 2, "images", "[A] AI images per carousel (rest HTML-rendered)"),
    ("gen_in", 10000, "tokens", "[A] Sonnet script generation input"), ("gen_out", 2500, "tokens", "[A]"),
    ("jd_in", 8000, "tokens", "[A] Sonnet script judge (mandatory)"), ("jd_out", 1000, "tokens", "[A]"),
    ("cap_in", 2000, "tokens", "[A] Haiku per-post caption/text"), ("cap_out", 400, "tokens", "[A]"),
    ("hj_in", 4000, "tokens", "[A] Haiku per-post judge (mandatory)"), ("hj_out", 500, "tokens", "[A]"),
    ("mix_th_full", 0.48, "share", "posting_plan_90d.csv lane mix D+7..D+90"), ("mix_mv_full", 0.26, "share", "same"),
    ("mix_in", 0.26, "share", "same (lean moves movement into talking-head)"),
    ("views_post", 6000, "views", "BLITZ.md §13.2 central mean views/post at page age 15 d"),
    ("views_growth", 1.6, "x/30 d", "BLITZ.md §13.2 (cap 40,000)"),
    ("plat_mult_sum", 3.5, "sum", "BLITZ.md §13.2 IG 1.3 + FB 1.0 + TT 0.7 + YT 0.5 (Threads/X counted at 0)"),
    ("mrr_price", 25, "$/mo", "BRIEF CANON UPDATE 2 founding price"),
    ("pay_pct", 0.029, "share", "Shopify Payments online 2.9% + 30c [verify plan]: shopify.com/blog/credit-card-processing-fees"),
    ("pay_fixed", 0.30, "$", "same"), ("refund", 0.12, "share", "BLITZ.md §13.1 refunds"),
]
V = {k: v for k, v, *_ in INPUTS}
SCEN = [(cfg, pages, q, amp) for cfg in ("full", "lean") for pages in (3, 5, 7) for q in (3, 6, 9) for amp in (0, 500, 1500)]


def per_master(cfg: str) -> dict:
    v = V
    th = v["img_th"] * v["rt_img"] * v["p_img"] + (v["sec"] + v["swap_s"]) * v["rt_ls"] * v["p_ls"] + v["ins_th"] * v["rt_ins"] * v["p_veo"]
    mv = (v["img_mv"] * v["rt_img"] * v["p_img"] + v["mv_share"] * v["sec"] * v["rt_mv"] * v["p_mv"]
          + ((1 - v["mv_share"]) * v["sec"] + v["swap_s"]) * v["rt_ls"] * v["p_ls"] + v["ins_mv"] * v["rt_ins"] * v["p_veo"])
    ins = v["img_in"] * v["rt_img"] * v["p_img"] + v["ins_in"] * v["rt_ins"] * v["p_veo"]
    mth, mmv = (v["mix_th_full"], v["mix_mv_full"]) if cfg == "full" else (v["mix_th_full"] + v["mix_mv_full"], 0)
    mix = {"th": mth, "mv": mmv, "in": v["mix_in"]}
    imgs = (mix["th"] * v["img_th"] + mix["mv"] * v["img_mv"] + mix["in"] * v["img_in"]) * v["rt_img"] + 2 * v["carousel_share"] * v["carousel_imgs"]
    lip = (mix["th"] * (v["sec"] + v["swap_s"]) + mix["mv"] * ((1 - v["mv_share"]) * v["sec"] + v["swap_s"])) * v["rt_ls"]
    mot = mix["mv"] * v["mv_share"] * v["sec"] * v["rt_mv"]
    veo = (mix["th"] * v["ins_th"] + mix["mv"] * v["ins_mv"] + mix["in"] * v["ins_in"]) * v["rt_ins"]
    chars = v["chars_s"] * v["sec"] * v["rt_tts"] + v["swap_chars"] * (mix["th"] + mix["mv"])
    claude = ((v["gen_in"] + v["jd_in"]) * v["p_sin"] + (v["gen_out"] + v["jd_out"]) * v["p_sout"]
              + 6 * ((v["cap_in"] + v["hj_in"]) * v["p_hin"] + (v["cap_out"] + v["hj_out"]) * v["p_hout"])) / 1e6
    return {"img_usd": imgs * v["p_img"], "lipsync_usd": lip * v["p_ls"], "motion_usd": mot * v["p_mv"], "veo_usd": veo * v["p_veo"],
            "chars": chars, "claude_usd": claude, "review_usd": 6 * v["review_s"] / 3600 * v["review_rate"],
            "lane_th": th, "lane_mv": mv, "lane_in": ins}


def eleven_month(chars: float) -> float:
    v = V
    return min(f + max(0, chars - inc) / 1000 * v["el_ovg"] for f, inc in
               [(v["el_creator_fee"], v["el_creator_inc"]), (v["el_pro_fee"], v["el_pro_inc"]),
                (v["el_scale_fee"], v["el_scale_inc"]), (v["el_bus_fee"], v["el_bus_inc"])])


def scenario(cfg, pages, q, amp) -> dict:
    v = V
    pm = per_master(cfg)
    m = pages * q                      # masters/day
    posts = m * 6
    files = m * 4
    d = {"config": cfg, "pages": pages, "posts_per_day_per_platform": q, "platforms": 6, "paid_amp_per_day": amp,
         "masters_per_day": m, "posts_per_day": posts}
    d["nano_banana"] = m * pm["img_usd"]
    d["kling_lipsync"] = m * pm["lipsync_usd"]
    d["kling_motion"] = m * pm["motion_usd"]
    d["veo_inserts"] = m * pm["veo_usd"]
    d["elevenlabs"] = eleven_month(m * pm["chars"] * DPM) / DPM
    d["claude_gen_judge"] = m * pm["claude_usd"]
    d["human_review"] = m * pm["review_usd"]
    d["performer_shoot"] = ((v["shoot_oneoff"] + v["shoot_refresh_q"]) / v["shoot_amort_days"]) if cfg == "full" else 0.0
    boxes = 1 + math.ceil(files * v["assembly_min_file"] / v["box_cpu_min_day"])
    d["vps_assembly"] = boxes * v["vps_box"] / DPM
    d["storage_cdn"] = (files * v["gb_file"] + m * v["gb_master_tmp"]) * v["retain_days"] * v["r2_gb"] / DPM
    d["fixed_saas"] = (v["supabase"] + v["vercel"] + v["n8n"] + v["email"]) / DPM
    d["posting_tools"] = (v["manychat_page"] * pages + v["posting_api"]) / DPM
    d["paid_amplification"] = float(amp)
    comps = ["nano_banana", "kling_lipsync", "kling_motion", "veo_inserts", "elevenlabs", "claude_gen_judge", "human_review",
             "performer_shoot", "vps_assembly", "storage_cdn", "fixed_saas", "posting_tools"]
    d["production_per_day"] = sum(d[c] for c in comps)
    d["total_per_day"] = d["production_per_day"] + amp
    d["total_per_month"] = d["total_per_day"] * DPM
    d["cost_per_post"] = d["production_per_day"] / posts
    views15 = m * v["views_post"] * v["plat_mult_sum"]
    views75 = views15 * v["views_growth"] ** 2
    d["organic_views_per_day_age15"] = views15
    d["cpm_views_age15"] = d["production_per_day"] / views15 * 1000
    d["cpm_views_age75"] = d["production_per_day"] / views75 * 1000
    net_member = (v["mrr_price"] * (1 - v["pay_pct"]) - v["pay_fixed"]) * (1 - v["refund"])
    d["breakeven_members"] = d["total_per_month"] / net_member
    d["breakeven_mrr"] = d["breakeven_members"] * v["mrr_price"]
    return d


def write_csv(rows):
    with OUT_CSV.open("w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=list(rows[0]))
        w.writeheader()
        for r in rows:
            w.writerow({k: (round(x, 4) if isinstance(x, float) else x) for k, x in r.items()})


# ------------------------------------------------------------------ the Costs sheet (formula-driven)
def write_sheet(rows):
    import openpyxl
    from openpyxl.styles import Font, PatternFill
    wb = openpyxl.load_workbook(XLSX)
    if "Costs" in wb.sheetnames:
        del wb["Costs"]
    ws = wb.create_sheet("Costs")
    bold = Font(bold=True)
    inp = PatternFill("solid", fgColor="FFF2CC")
    ws["A1"] = "Costs: content production at 3/6/9 posts/day/platform x 3/5/7 pages x 6 platforms (edit yellow inputs; everything below is formulas)"
    ws["A1"].font = bold
    ws.append([])
    ws.append(["Input", "Value", "Unit", "Source / label ([A] = assumption)"])
    for c in ws[3]:
        c.font = bold
    ref = {}
    for k, val, unit, src in INPUTS:
        ws.append([k, val, unit, src])
        r = ws.max_row
        ws.cell(r, 2).fill = inp
        ref[k] = f"$B${r}"
    ws.append(["days_per_month", DPM, "days", "30.4"])
    ref["dpm"] = f"$B${ws.max_row}"
    R = ref
    # derived per-master block
    ws.append([])
    ws.append(["Per master (one source script -> 4 video files + 2 text posts)", "full", "lean"])
    for c in ws[ws.max_row]:
        c.font = bold
    def mixes(col):
        th = f"({R['mix_th_full']}+{R['mix_mv_full']})" if col == "lean" else R["mix_th_full"]
        mv = "0" if col == "lean" else R["mix_mv_full"]
        return th, mv, R["mix_in"]
    pm_rows = {}
    def pm(name, f_full, f_lean):
        ws.append([name, "=" + f_full, "=" + f_lean])
        pm_rows[name] = ws.max_row
    def build(col):
        th, mv, mi = mixes(col)
        return {
            "images": f"(({th})*{R['img_th']}+({mv})*{R['img_mv']}+{mi}*{R['img_in']})*{R['rt_img']}+2*{R['carousel_share']}*{R['carousel_imgs']}",
            "lipsync_s": f"(({th})*({R['sec']}+{R['swap_s']})+({mv})*((1-{R['mv_share']})*{R['sec']}+{R['swap_s']}))*{R['rt_ls']}",
            "motion_s": f"({mv})*{R['mv_share']}*{R['sec']}*{R['rt_mv']}",
            "veo_s": f"(({th})*{R['ins_th']}+({mv})*{R['ins_mv']}+{mi}*{R['ins_in']})*{R['rt_ins']}",
            "tts_chars": f"{R['chars_s']}*{R['sec']}*{R['rt_tts']}+{R['swap_chars']}*(({th})+({mv}))",
            "claude_usd": f"(({R['gen_in']}+{R['jd_in']})*{R['p_sin']}+({R['gen_out']}+{R['jd_out']})*{R['p_sout']}+6*(({R['cap_in']}+{R['hj_in']})*{R['p_hin']}+({R['cap_out']}+{R['hj_out']})*{R['p_hout']}))/1000000",
            "review_usd": f"6*{R['review_s']}/3600*{R['review_rate']}",
        }
    bf, bl = build("full"), build("lean")
    for k in bf:
        pm(k, bf[k], bl[k])
    P = {k: (f"$B${r}", f"$C${r}") for k, r in pm_rows.items()}
    ws.append([])
    hdr = ["config", "pages", "posts/day/platform", "paid amp $/day", "masters/day", "posts/day", "Nano Banana", "Kling lip-sync",
           "Kling motion", "Veo inserts", "ElevenLabs", "Claude gen+judge", "Human review", "Performer shoot", "VPS assembly",
           "Storage/CDN", "n8n/Supabase/Vercel/email", "Posting tools", "Paid amplification", "Production $/day", "Total $/day",
           "Total $/month", "$/post", "$/1K views (age 15)", "$/1K views (age 75)", "Breakeven members", "Breakeven MRR"]
    ws.append(hdr)
    hr = ws.max_row
    for c in ws[hr]:
        c.font = bold
    for cfg, pages, q, amp in SCEN:
        r = ws.max_row + 1
        i = 0 if cfg == "full" else 1
        g = lambda k: P[k][i]  # noqa: E731
        el = (f"MIN({R['el_creator_fee']}+MAX(0,E{r}*{g('tts_chars')}*{R['dpm']}-{R['el_creator_inc']})/1000*{R['el_ovg']},"
              f"{R['el_pro_fee']}+MAX(0,E{r}*{g('tts_chars')}*{R['dpm']}-{R['el_pro_inc']})/1000*{R['el_ovg']},"
              f"{R['el_scale_fee']}+MAX(0,E{r}*{g('tts_chars')}*{R['dpm']}-{R['el_scale_inc']})/1000*{R['el_ovg']},"
              f"{R['el_bus_fee']}+MAX(0,E{r}*{g('tts_chars')}*{R['dpm']}-{R['el_bus_inc']})/1000*{R['el_ovg']})/{R['dpm']}")
        ws.append([
            cfg, pages, q, amp, f"=B{r}*C{r}", f"=E{r}*6",
            f"=E{r}*{g('images')}*{R['p_img']}", f"=E{r}*{g('lipsync_s')}*{R['p_ls']}", f"=E{r}*{g('motion_s')}*{R['p_mv']}",
            f"=E{r}*{g('veo_s')}*{R['p_veo']}", "=" + el, f"=E{r}*{g('claude_usd')}", f"=E{r}*{g('review_usd')}",
            f'=IF(A{r}="full",({R["shoot_oneoff"]}+{R["shoot_refresh_q"]})/{R["shoot_amort_days"]},0)',
            f"=(1+ROUNDUP(E{r}*4*{R['assembly_min_file']}/{R['box_cpu_min_day']},0))*{R['vps_box']}/{R['dpm']}",
            f"=(E{r}*4*{R['gb_file']}+E{r}*{R['gb_master_tmp']})*{R['retain_days']}*{R['r2_gb']}/{R['dpm']}",
            f"=({R['supabase']}+{R['vercel']}+{R['n8n']}+{R['email']})/{R['dpm']}",
            f"=({R['manychat_page']}*B{r}+{R['posting_api']})/{R['dpm']}",
            f"=D{r}", f"=SUM(G{r}:R{r})", f"=T{r}+S{r}", f"=U{r}*{R['dpm']}", f"=T{r}/F{r}",
            f"=T{r}/(E{r}*{R['views_post']}*{R['plat_mult_sum']})*1000",
            f"=T{r}/(E{r}*{R['views_post']}*{R['plat_mult_sum']}*{R['views_growth']}^2)*1000",
            f"=V{r}/(({R['mrr_price']}*(1-{R['pay_pct']})-{R['pay_fixed']})*(1-{R['refund']}))",
            f"=Z{r}*{R['mrr_price']}",
        ])
        for col in range(7, 28):
            ws.cell(r, col).number_format = "#,##0.00"
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["D"].width = 60
    wb.save(XLSX)
    return hr


if __name__ == "__main__":
    lm = ROOT / "data/content/posting_plan_lane_mix.json"
    if lm.exists():
        mix = json.loads(lm.read_text())["lane_mix_d7_plus"]
        V["mix_th_full"], V["mix_mv_full"], V["mix_in"] = (round(mix.get("talking_head", 0), 2), round(mix.get("movement", 0), 2),
                                                           round(mix.get("insert", 0), 2))
        INPUTS[:] = [(k, V[k], u, s) for k, _, u, s in INPUTS]
    rows = [scenario(*s) for s in SCEN]
    write_csv(rows)
    if "--no-xlsx" not in sys.argv:
        write_sheet(rows)
    print(json.dumps({k: round(v, 3) for k, v in per_master("full").items()}), json.dumps({k: round(v, 3) for k, v in per_master("lean").items()}))
    for r in rows:
        if r["paid_amp_per_day"] == 0:
            print(f"{r['config']:4} p{r['pages']} q{r['posts_per_day_per_platform']} posts/d {r['posts_per_day']:4} $/d {r['total_per_day']:8.0f} $/mo {r['total_per_month']:9.0f} $/post {r['cost_per_post']:.2f} cpm15 {r['cpm_views_age15']:.3f} be_mrr {r['breakeven_mrr']:8.0f}")
